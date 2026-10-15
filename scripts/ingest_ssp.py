#!/usr/bin/env python3
"""Baixa dados criminais oficiais da SSP e insere linhas novas no PostGIS."""

import argparse
import json
import math
import tempfile
from datetime import date, datetime, time
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
import psycopg
from dotenv import dotenv_values
from openpyxl import load_workbook
from psycopg import sql

ROOT = Path(__file__).resolve().parents[1]
PORTAL = "https://www.ssp.sp.gov.br/estatistica/consultas"
CATALOG = (
    "https://www.ssp.sp.gov.br/assets/estatistica/transparencia/spDados160_516.json"
)
COLUMNS = """nome_departamento nome_seccional nome_delegacia nome_municipio num_bo
ano_bo data_registro data_ocorrencia_bo hora_ocorrencia_bo desc_periodo descr_tipolocal
 descr_subtipolocal bairro logradouro numero_logradouro latitude longitude
nome_delegacia_circunscricao nome_departamento_circunscricao nome_seccional_circunscricao
nome_municipio_circunscricao rubrica descr_conduta natureza_apurada mes_estatistica
ano_estatistica cmd btl cia cod_ibge""".split()
INTEGERS = {"ano_bo", "mes_estatistica", "ano_estatistica", "cod_ibge"}
DATES = {"data_registro", "data_ocorrencia_bo"}


def header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def normalize(column, value):
    if value is None or (
        isinstance(value, str) and value.strip().upper() in {"", "NULL", "-"}
    ):
        return None
    if column in DATES:
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(str(value).strip(), fmt).date()
            except ValueError:
                pass
        raise ValueError(f"data inválida em {column}")
    if column in INTEGERS:
        number = float(str(value).replace(",", "."))
        if not math.isfinite(number) or not number.is_integer():
            raise ValueError(f"inteiro inválido em {column}")
        number = int(number)
        if column == "mes_estatistica" and not 1 <= number <= 12:
            raise ValueError("mês inválido")
        if column != "cod_ibge" and not -32768 <= number <= 32767:
            raise ValueError(f"inteiro fora do limite em {column}")
        return number
    if column in {"latitude", "longitude"}:
        number = float(str(value).strip().replace(",", "."))
        maximum = 90 if column == "latitude" else 180
        if not math.isfinite(number) or abs(number) > maximum:
            raise ValueError(f"coordenada inválida em {column}")
        return number
    if isinstance(value, time):
        return value.isoformat()
    return str(value).strip()


def official_url(url):
    parsed = urlparse(url)
    if parsed.scheme != "https" or not (parsed.hostname or "").endswith(
        ".ssp.sp.gov.br"
    ):
        raise ValueError("Use um link HTTPS oficial da SSP (*.ssp.sp.gov.br).")
    return url


def resolve_url(client, url, year):
    official_url(url)
    if urlparse(url).path.lower().endswith(".xlsx"):
        return url
    if urlparse(url).path.rstrip("/") != "/estatistica/consultas":
        raise ValueError(
            "Informe a página /estatistica/consultas ou o link de uma planilha XLSX."
        )
    if year is None:
        raise ValueError("--ano é obrigatório ao usar a página de consultas.")
    response = client.get(CATALOG)
    response.raise_for_status()
    catalog = json.loads(response.content.decode("utf-8-sig"))
    for group in catalog["data"]:
        if group["nome"] == "Dados criminais":
            for entry in group["lista"]:
                if str(entry["periodo"]) == str(year):
                    return official_url(
                        urljoin(CATALOG, "/" + entry["arquivo"].lstrip("/"))
                    )
    raise ValueError(f"Ano {year} não encontrado no catálogo da SSP.")


def rows(path, municipality=None, limit=None):
    workbook = load_workbook(path, read_only=True, data_only=True)
    count = 0
    invalid_coordinates = 0
    found = False
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            names = [header(value) for value in next(iterator, ())]
            if "num_bo" not in names or "rubrica" not in names:
                continue  # Ignora a aba de documentação da SSP.
            missing = set(COLUMNS) - set(names)
            if missing:
                raise ValueError(f"{sheet.title}: colunas ausentes: {sorted(missing)}")
            found = True
            positions = [names.index(column) for column in COLUMNS]
            for line, raw in enumerate(iterator, 2):
                if not any(raw):
                    continue
                if municipality is not None:
                    raw_ibge = raw[names.index("cod_ibge")]
                    if normalize("cod_ibge", raw_ibge) != municipality:
                        continue
                try:
                    values = []
                    bad_coordinates = False
                    for column, index in zip(COLUMNS, positions, strict=True):
                        try:
                            values.append(normalize(column, raw[index]))
                        except (ValueError, TypeError, OverflowError):
                            if column not in {"latitude", "longitude"}:
                                raise
                            values.append(None)
                            bad_coordinates = True
                    if bad_coordinates:
                        invalid_coordinates += 1
                        values[COLUMNS.index("latitude")] = None
                        values[COLUMNS.index("longitude")] = None
                    record = tuple(values)
                except (ValueError, TypeError, OverflowError) as exc:
                    raise ValueError(f"{sheet.title}, linha {line}: {exc}") from exc
                yield record
                count += 1
                if limit is not None and count >= limit:
                    return
        if not found:
            raise ValueError("Nenhuma aba de dados criminais compatível encontrada.")
    finally:
        workbook.close()
        if invalid_coordinates:
            print(
                f"Coordenadas inválidas em {invalid_coordinates} linhas; pontos não gerados.",
                flush=True,
            )


def connect(target):
    values = dotenv_values(ROOT / "deploy" / f".env.{target}")
    expected = f"rotasegura_{target}"
    if values.get("POSTGRES_DATABASE") != expected:
        raise ValueError(f"O arquivo do ambiente deve apontar para {expected}.")
    return psycopg.connect(
        host="127.0.0.1",
        port=int(values["DATABASE_PORT"]),
        dbname=expected,
        user=values["POSTGRES_ADMIN_USER"],
        password=values["POSTGRES_ADMIN_PASSWORD"],
        connect_timeout=10,
    )


def ingest(connection, records, dry_run=False):
    columns = sql.SQL(", ").join(map(sql.Identifier, COLUMNS))
    with connection.transaction():
        with connection.cursor() as cursor:
            # Serializa execuções deste importador para evitar inserções simultâneas repetidas.
            cursor.execute(
                "SELECT pg_advisory_xact_lock(hashtext('ssp_occurrences_ingestion'))"
            )
            cursor.execute(
                sql.SQL(
                    "CREATE TEMP TABLE ssp_load ON COMMIT DROP AS SELECT {} FROM public.ocorrencias_stg WITH NO DATA"
                ).format(columns)
            )
            total = 0
            with cursor.copy(
                sql.SQL("COPY ssp_load ({}) FROM STDIN").format(columns)
            ) as copy:
                for record in records:
                    copy.write_row(record)
                    total += 1
                    if total % 10000 == 0:
                        print(f"Validadas {total} linhas...", flush=True)
            cursor.execute(
                sql.SQL("""CREATE TEMP TABLE ssp_new ON COMMIT DROP AS
                SELECT {columns} FROM ssp_load EXCEPT SELECT {columns} FROM public.ocorrencias_stg""").format(
                    columns=columns
                )
            )
            cursor.execute("SELECT count(*) FROM ssp_new")
            new = cursor.fetchone()[0]
            if not dry_run:
                # id é gerado pelo banco. Zero é usado pela SSP para localização indisponível.
                cursor.execute(
                    sql.SQL("""INSERT INTO public.ocorrencias_stg ({columns}, bairro_normalizado, geom, geog)
                    SELECT {columns}, lower(translate(bairro,
                    'ÁÀÂÃÄÉÈÊËÍÌÎÏÓÒÔÕÖÚÙÛÜÇáàâãäéèêëíìîïóòôõöúùûüç',
                    'AAAAAEEEEIIIIOOOOOUUUUCaaaaaeeeeiiiiooooouuuuc')),
                    CASE WHEN latitude IS NOT NULL AND longitude IS NOT NULL AND latitude <> 0 AND longitude <> 0
                         THEN ST_SetSRID(ST_MakePoint(longitude, latitude), 4326) END,
                    CASE WHEN latitude IS NOT NULL AND longitude IS NOT NULL AND latitude <> 0 AND longitude <> 0
                         THEN ST_SetSRID(ST_MakePoint(longitude, latitude), 4326)::geography END
                    FROM ssp_new""").format(columns=columns)
                )
                cursor.execute("ANALYZE public.ocorrencias_stg")
            cursor.execute("DROP TABLE ssp_new, ssp_load")
    return {
        "lidas": total,
        "novas": new,
        "duplicadas": total - new,
        "inseridas": 0 if dry_run else new,
        "simulacao": dry_run,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--url",
        default=PORTAL,
        help="Página de consultas ou link HTTPS de uma planilha XLSX da SSP",
    )
    parser.add_argument("--ano", type=int)
    parser.add_argument("--target", choices=["test", "dev"], default="test")
    parser.add_argument(
        "--municipio-ibge", type=int, help="Exemplo: 3550308 para São Paulo"
    )
    parser.add_argument(
        "--limit", type=int, help="Limita a amostra; permitido somente em test"
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="Valida e conta sem inserir"
    )
    args = parser.parse_args()
    if args.limit is not None and (args.target != "test" or args.limit < 1):
        parser.error("--limit deve ser positivo e só pode ser usado com --target test")
    try:
        with httpx.Client(
            timeout=httpx.Timeout(120, connect=30), follow_redirects=False
        ) as client:
            url = resolve_url(client, args.url, args.ano)
            print(f"Fonte: {url}\nDestino: rotasegura_{args.target}", flush=True)
            with tempfile.TemporaryDirectory(prefix="ssp-ingest-") as directory:
                path = Path(directory) / "ssp.xlsx"
                with client.stream("GET", url) as response:
                    response.raise_for_status()
                    with path.open("wb") as output:
                        for chunk in response.iter_bytes():
                            output.write(chunk)
                print(f"Download concluído ({path.stat().st_size} bytes).", flush=True)
                with connect(args.target) as connection:
                    report = ingest(
                        connection,
                        rows(path, args.municipio_ibge, args.limit),
                        args.dry_run,
                    )
                print(json.dumps(report, ensure_ascii=False))
    except (ValueError, KeyError, httpx.HTTPError, psycopg.Error) as exc:
        parser.exit(1, f"Falha na ingestão: {exc}\n")


if __name__ == "__main__":
    main()
